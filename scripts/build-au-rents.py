#!/usr/bin/env python3
"""Builds lib/data/australiaRents.json — official local rents (and NSW
sale prices) from the states' own bond / transfer records:
 - NSW: Department of Communities and Justice "Rent and Sales Report"
   (dcj.nsw.gov.au): Table 2 weekly rents of NEW BONDS per postcode ×
   dwelling type × bedrooms (quarterly); Table 4 sale prices per postcode,
   strata (units, townhouses, villas) vs non-strata (houses), from the
   'Notice of Sale or Transfer of Land' forms lodged with NSW Land
   Registry; areas with 10 or fewer sales are not published, "s" = 11–30.
 - SA: "Private Rental Report" (data.sa.gov.au, quarterly): weekly rents
   of bonds lodged per suburb × flats / houses × bedrooms (Consumer and
   Business Services bond data); counts rounded to 5, "*" = 1–5.
 - suburb → postcode (NSW) from ABS ASGS 2021 allocation files (mesh
   blocks: SAL × POA), only when 80%+ of the suburb's area is in one
   postcode.
A figure is used only with 10+ bonds / sales.

  python3 scripts/build-au-rents.py      (quarterly)
"""
import collections, io, json, pathlib, re, urllib.request
import openpyxl

UA = {"User-Agent": "Mozilla/5.0 (Pradixium/1.0)"}
TMP = pathlib.Path("/tmp/claude-0"); TMP.mkdir(parents=True, exist_ok=True)
def get(url):
    return urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=300).read()
def sheet(b, name):
    return list(openpyxl.load_workbook(io.BytesIO(b), read_only=True)[name].iter_rows(values_only=True))
def num(x):
    try: return float(str(x).replace(",", ""))
    except (TypeError, ValueError): return None

# ---- NSW: the report page lists the latest rent and sales tables
NSW_PAGE = "https://dcj.nsw.gov.au/about-us/families-and-communities-statistics/housing-rent-and-sales/rent-and-sales-report.html"
page = get(NSW_PAGE).decode("utf-8", "ignore")
links = re.findall(r'href="([^"]*/(rent|sales)-tables-([a-z]+)-(\d{4})-quarter\.xlsx)"', page)
latest = {}
for href, kind, mon, yr in links:
    latest.setdefault(kind, ("https://dcj.nsw.gov.au" + href if href.startswith("/") else href, f"{mon.title()} quarter {yr}"))
out = {"nsw": {}, "sa": {}}

rows = sheet(get(latest["rent"][0]), "Postcode")
period = next(str(r[0]) for r in rows if r and r[0] and str(r[0]).startswith("Reporting period")).replace("Reporting period: ", "")
hdr = next(i for i, r in enumerate(rows) if r and r[0] == "Postcode")
rent = collections.defaultdict(dict)
for r in rows[hdr + 1:]:
    if not r or not isinstance(r[0], (int, float)): continue
    pc, typ, beds, med, n = str(int(r[0])).zfill(4), str(r[1]), str(r[2]), num(r[4]), r[6]
    nn = num(n)
    if med is None or (nn is None and n != "s") or (nn is not None and nn < 10): continue
    rent[pc][f"{typ}|{beds}"] = [med, int(nn) if nn else "11–30"]
out["nsw"]["rent"] = {"period": period, "source": "NSW Department of Communities and Justice — Rent and Sales Report, Table 2 (weekly rents of new bonds by postcode)", "sourceUrl": NSW_PAGE, "pc": rent}

rows = sheet(get(latest["sales"][0]), "Postcode")
speriod = next(str(r[0]) for r in rows if r and r[0] and str(r[0]).startswith("Reporting period")).replace("Reporting period: ", "")
hdr = next(i for i, r in enumerate(rows) if r and r[0] == "Postcode")
sales = collections.defaultdict(dict)
for r in rows[hdr + 1:]:
    if not r or not isinstance(r[0], (int, float)): continue
    pc, typ = str(int(r[0])).zfill(4), str(r[1])
    q1, med, q3, n = num(r[2]), num(r[3]), num(r[4]), r[6]
    if med is None or typ not in ("Strata", "Non Strata"): continue
    nn = num(n)
    sales[pc][typ] = [round(med * 1000), round(q1 * 1000) if q1 else None, round(q3 * 1000) if q3 else None, int(nn) if nn else "11–30"]
out["nsw"]["sales"] = {"period": speriod, "source": "NSW Department of Communities and Justice — Rent and Sales Report, Table 4 (sale prices by postcode, from Notices of Sale lodged with NSW Land Registry)", "sourceUrl": NSW_PAGE, "pc": sales}

# ---- SA private rental report (latest XLSX on data.sa.gov.au)
pkg = json.loads(get("https://data.sa.gov.au/data/api/3/action/package_show?id=private-rent-report"))["result"]
res = sorted([x for x in pkg["resources"] if x["format"].upper() == "XLSX"], key=lambda x: x.get("created") or "")[-1]
rows = sheet(get(res["url"]), "Suburb")
title = next(str(r[0]) for r in rows if r and r[0] and "median weekly rental" in str(r[0]))
saperiod = re.search(r"rental, (.+?), South Australia", title).group(1)
# columns: flats 1-4+ (count, median) at 1..8, flats total 9/10; houses 11..18, total 19/20
COLS = [("Flat", b, 1 + 2 * i) for i, b in enumerate(("1", "2", "3", "4+"))] + [("Flat", "Total", 9)] + \
       [("House", b, 11 + 2 * i) for i, b in enumerate(("1", "2", "3", "4+"))] + [("House", "Total", 19)]
start = next(i for i, r in enumerate(rows) if r and r[0] == "Row Labels")
sa = {}
for r in rows[start + 1:]:
    if not r or not r[0] or r[0] in ("Metro", "Country", "Grand Total") or r[1:] == (None,) * (len(r) - 1): continue
    d = {}
    for typ, b, c in COLS:
        n, med = num(r[c]), num(r[c + 1])
        if n is not None and n >= 10 and med: d[f"{typ}|{b}"] = [med, int(n)]
    if d: sa.setdefault(str(r[0]).strip().upper(), d)
out["sa"]["rent"] = {"period": saperiod, "source": "Government of South Australia — Private Rental Report (" + res["name"] + "), bonds lodged with Consumer and Business Services", "sourceUrl": "https://data.sa.gov.au/data/dataset/private-rent-report", "sub": sa}

# ---- NSW suburb → postcode (ABS mesh blocks)
ALLOC = "https://www.abs.gov.au/statistics/standards/australian-statistical-geography-standard-asgs/edition-3-july-2021-june-2026/access-and-downloads/allocation-files/"
def rows_x(name):
    p = TMP / name
    if not p.exists(): p.write_bytes(get(ALLOC + name))
    it = openpyxl.load_workbook(p, read_only=True).worksheets[0].iter_rows(values_only=True)
    h = next(it)
    for r in it: yield dict(zip(h, r))
poa = {r["MB_CODE_2021"]: r["POA_CODE_2021"] for r in rows_x("POA_2021_AUST.xlsx")}
area = collections.defaultdict(collections.Counter)
for r in rows_x("SAL_2021_AUST.xlsx"):
    if r["STATE_NAME_2021"] != "New South Wales" or r["MB_CODE_2021"] not in poa: continue
    area[r["SAL_NAME_2021"]][poa[r["MB_CODE_2021"]]] += float(r["AREA_ALBERS_SQKM"] or 0)
sub_pc = {}
for name, c in area.items():
    tot = sum(c.values()); pc, a = c.most_common(1)[0]
    if tot and a / tot >= 0.8 and re.match(r"^\d{4}$", str(pc)):
        sub_pc[re.sub(r"\s*\([^)]*\)$", "", name).strip().upper()] = pc
out["nsw"]["suburbPostcode"] = sub_pc

p = pathlib.Path(__file__).resolve().parent.parent / "lib/data/australiaRents.json"
p.write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")))
print("NSW rent", period, len(rent), "postcodes; sales", speriod, len(sales), "postcodes;", len(sub_pc), "suburbs → postcode")
print("SA rent", saperiod, len(sa), "suburbs")
