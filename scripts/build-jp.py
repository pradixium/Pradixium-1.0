#!/usr/bin/env python3
"""Japan → lib/data/japanPrices.json (monthly, after MLIT's release)

MLIT 不動産価格指数（住宅） (Real Estate Price Index, residential; 2010
average = 100; from registered transactions): the latest month of the
ORIGINAL series (原系列) for every sheet MLIT publishes — Japan, the 9
regions, the 3 metro areas, Tokyo / Aichi / Osaka prefectures — by type:
all residential, land, detached houses, condominiums (区分所有); change on
the same month a year earlier + MLIT's sample count.
  python3 scripts/build-jp.py
"""
import io, json, os, re, urllib.request
import openpyxl
UA = {"User-Agent": "Mozilla/5.0", "Accept": "*/*"}
PAGE = "https://www.mlit.go.jp/totikensangyo/totikensangyo_tk5_000085.html"
page = urllib.request.urlopen(urllib.request.Request(PAGE, headers=UA), timeout=60).read().decode("utf-8", "ignore")
# the first Excel under 最新データ is 不動産価格指数（住宅）
href = re.search(r'最新データ.*?不動産価格指数（住宅）.*?href="([^"]+\.xlsx)"', page, re.S).group(1)
url = "https://www.mlit.go.jp" + href if href.startswith("/") else href
wb = openpyxl.load_workbook(io.BytesIO(urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=120).read()), read_only=True, data_only=True)
COLS = {"all": (1, 3), "land": (4, 6), "detached": (7, 9), "condo": (10, 12)}
out = {"source": "MLIT — 不動産価格指数（住宅） Real Estate Price Index (residential), original series, 2010 = 100", "sourceUrl": PAGE, "file": url, "areas": {}}
for name in wb.sheetnames:
    if not name.endswith("原系列"): continue
    rows = [r for r in wb[name].iter_rows(values_only=True) if r and hasattr(r[0], "year")]
    if not rows: continue   # the "----->>原系列" separator sheet
    last = rows[-1]
    ago = next(r for r in rows if r[0].year == last[0].year - 1 and r[0].month == last[0].month)
    area = {}
    for k, (vi, ni) in COLS.items():
        if isinstance(last[vi], (int, float)) and isinstance(ago[vi], (int, float)):
            area[k] = {"index": round(last[vi], 2), "yearAgo": round(ago[vi], 2), "yoyPercent": round((last[vi] / ago[vi] - 1) * 100, 1), "sample": last[ni]}
    out["areas"][name.replace("原系列", "")] = area
    out["period"] = f"{last[0].year}-{last[0].month:02d}"
p = os.path.join(os.path.dirname(__file__), "..", "lib", "data", "japanPrices.json")
json.dump(out, open(p, "w"), ensure_ascii=False, indent=1)
print(out["period"], len(out["areas"]), {k: v.get("condo", {}).get("yoyPercent") for k, v in out["areas"].items()})
