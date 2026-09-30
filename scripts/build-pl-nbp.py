#!/usr/bin/env python3
"""Poland: Narodowy Bank Polski — BaRN house prices database, TRANSACTION
prices (not offer prices) per m² of flats in the 16 voivodeship capitals and
Gdynia, secondary and primary market, latest quarter (average price, VAT
included; data reported to the NBP by estate agents and developers under the
official statistics programme). → lib/data/polandNbpPrices.json
Re-run each quarter:  pip install openpyxl && python3 scripts/build-pl-nbp.py
"""
import io, json, os, urllib.request
import openpyxl

URL = "https://static.nbp.pl/dane/rynek-nieruchomosci/ceny_mieszkan.xlsx"
wb = openpyxl.load_workbook(io.BytesIO(urllib.request.urlopen(urllib.request.Request(URL, headers={"User-Agent": "Mozilla/5.0 (Pradixium data build)"}), timeout=120).read()), read_only=True, data_only=True)

def block(sheet):
    rows = list(wb[sheet].iter_rows(values_only=True))
    hdr = rows[6]
    for k, c in enumerate(hdr):
        if c != "Kwartał": continue
        label = " ".join(str(rows[r][k + 1] or "") + " " + str(rows[r][k] or "") for r in range(0, 6))
        if "transakcyjne" not in label: continue
        cities = []
        for j in range(k + 1, len(hdr)):
            if not hdr[j] or hdr[j] == "Kwartał": break
            cities.append((j, str(hdr[j]).rstrip("*")))
        last = None
        for r in rows[7:]:
            q = r[k]
            if isinstance(q, str) and q.strip() and q.split()[-1].isdigit(): last = r
        return last[k], {name: round(last[j]) for j, name in cities if isinstance(last[j], (int, float)) and not name[0].isdigit()}
    raise SystemExit(f"no transaction block in {sheet}")

def hedonic_yoy():
    rows = list(wb["Rynek wtórny"].iter_rows(values_only=True))
    hdr = rows[6]
    for k, c in enumerate(hdr):
        if c != "Kwartał": continue
        label = " ".join(str(rows[r][k + 1] or "") for r in range(0, 6))
        if "y-o-y" not in label and "analogicznego" not in label: continue
        cities = []
        for j in range(k + 1, len(hdr)):
            if not hdr[j] or hdr[j] == "Kwartał" or hdr[j] == "Indeks": break
            cities.append((j, str(hdr[j]).rstrip("*")))
        # this block leaves its quarter column empty: the row's quarter is the
        # sheet's first column (same rows)
        last = None
        for r in rows[7:]:
            q = r[0]
            if isinstance(q, str) and q.strip() and q.split()[-1].isdigit() and any(isinstance(r[j], (int, float)) for j, _ in cities): last = r
        if last is None: return None, {}
        out = {}
        for j, name in cities:
            v = last[j]
            if not isinstance(v, (int, float)) or name[0].isdigit(): continue
            for city in (["Gdańsk", "Gdynia", "Sopot"] if name == "Trójmiasto" else [name]): out[city] = round(v - 100, 1)
        return last[0], out
    return None, {}

rq, resale = block("Rynek wtórny")
nq, new = block("Rynek pierwotny")
assert rq == nq, (rq, nq)
rom, year = rq.split()
out = {"source": "Narodowy Bank Polski (NBP) — BaRN house prices database, transaction prices",
       "sourceUrl": "https://nbp.pl/publikacje/cykliczne-materialy-analityczne-nbp/rynek-nieruchomosci/informacja-kwartalna/",
       "file": URL, "quarter": f"Q{ {'I': 1, 'II': 2, 'III': 3, 'IV': 4}[rom] } {year}",
       "resale": resale, "new": new}
hq, yoy = hedonic_yoy()
if hq == rq:
    out["hedonicYoY"] = yoy   # NBP hedonic index, resale flats, change on the same quarter a year earlier (quality-adjusted)
assert out["resale"].get("Kraków") and out["resale"].get("Warszawa"), "format changed"
dst = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "lib", "data", "polandNbpPrices.json")
json.dump(out, open(dst, "w"), ensure_ascii=False, indent=1)
print(out["quarter"], "resale", out["resale"], "\nnew", out["new"], "\nyoy", out.get("hedonicYoY"))
